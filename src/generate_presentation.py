import os
import sys
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation(output_path="docs/Presentation.pptx"):
    print("Generating Presentation.pptx...")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    metrics_df = pd.read_csv("outputs/metrics.csv")
    bi_df = pd.read_csv("outputs/bi_summary.csv")
    adv_df = pd.read_csv("outputs/adversarial_results.csv")
    m = metrics_df.iloc[0]
    bi = bi_df.iloc[0]

    def add_header(slide, title, category="EXPLAINABLE AML PIPELINE | 22MDC56"):
        # Header banner
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        p_cat = tf.paragraphs[0]
        p_cat.text = category
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(59, 130, 246)

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(15, 23, 42)

    def set_speaker_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    txBox = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "Explainable ML Pipeline for AML Fraud Detection & Business Risk Management"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(30, 58, 138)

    p2 = tf.add_paragraph()
    p2.text = "\nCourse: 22MDC56 Machine Learning Laboratory | Academic Year 2025–2026"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(71, 85, 105)

    p3 = tf.add_paragraph()
    p3.text = "\nProject Team:\n• Sameul D (2403717673721040)\n• Mounica K (2403717673722028)\n• Naveen raj (2403717673721033)"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(51, 65, 85)
    set_speaker_notes(s1, "Welcome examiners. Today we present our end-to-end explainable ML AML pipeline connecting detection to business-risk decisions.")

    # Slide 2: Problem Statement
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Real-World AML Crisis & Problem Statement")
    tx = s2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Operational & Regulatory Dilemma in Modern Banking:"
    p.font.size = Pt(16)
    p.font.bold = True
    bullets = [
        "Massive Transaction Velocity: Global institutions handle hundreds of millions of daily transactions, with true illicit transfers < 0.1%.",
        "Severe Alert Fatigue: Traditional static rules generate >95% false alarms, consuming millions in compliance investigator salaries.",
        "Deterministic Evasion: Criminal syndicates structure amounts under fixed thresholds (smurfing, fan-out/fan-in) bypassing legacy filters.",
        "Regulatory Transparency Gap: Regulators mandate explainable audit trails (FATF, EU AMLD) that standard black-box deep models cannot satisfy.",
        "Core Project Objective: Build a zero-leakage, explainable ensemble ML detector with a cost-optimal BI layer that minimizes unrecovered fraud loss."
    ]
    for b in bullets:
        p_b = tf.add_paragraph()
        p_b.text = f"• {b}"
        p_b.font.size = Pt(14)
    set_speaker_notes(s2, "Highlight the extreme 0.1% rarity and explain how traditional rule filters cause 95% false positive rates.")

    # Slide 3: Literature Survey Table
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Comprehensive Literature Survey: Context & Prior Work")
    tx = s3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    lit_points = [
        "[1] Choi Min-Jae (2026): Machine learning in fraud detection and risk governance. Proved standard accuracy is deceptive in class imbalance and recommended PR-AUC; however, stopped short of modeling transaction dollar loss.",
        "[2] S. K. Bhowmick et al. (2026): 12-month core banking deployment using hybrid CNN/LSTM/Autoencoder + RF/GB. Reduced FPR from 32% to 2%, but PGD adversarial attack dropped accuracy to 71.2%, exposing acute evasion vulnerability.",
        "[3] S. Malathy et al. (2026): Multi-layered cybersecurity (SECURE framework) for high-frequency trading achieving 99.8% nominal accuracy; relied on non-temporal splits susceptible to look-ahead leakage.",
        "[4] S. Papastergiou et al. (2026): Integrated CodeBERT with SHAP and NIST 800-53 controls to map risk scores into cyber insurability underwriting premiums, demonstrating explainability as an actionable business driver.",
        "[5] R. Aljaafreh et al. (2026): Systematic review showing only 8.6% of modern ML implementations possess real-time integration with Business Intelligence (BI) layers, confirming a pervasive research void."
    ]
    for lp in lit_points:
        p_l = tf.add_paragraph()
        p_l.text = f"• {lp}"
        p_l.font.size = Pt(12)
    set_speaker_notes(s3, "Detail the five peer-reviewed 2026 papers, highlighting Bhowmick's finding that adversarial evasion dropped accuracy to 71.2%.")

    # Slide 4: Gap Analysis
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Research Gap Analysis & Fundamental Innovations")
    tx = s4.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    gaps = [
        "Gap 1: Disconnection from Financial Business Losses — Academic models treat a $10 error and a $1,000,000 error identically. Our pipeline integrates an asymmetric cost matrix ($FN = Amount, FP = $50$).",
        "Gap 2: Look-Ahead Temporal Leakage — Standard k-fold splitting contaminates training with future transactions. We enforce strict chronological partitioning ($max(train) < min(val) < min(test)$).",
        "Gap 3: Fragility to Adversarial Structuring — Attackers actively tune transaction features to evade detection. We subject models to black-box evasion tests and implement adversarial defense retraining.",
        "Gap 4: Inactionable Black-Box Outputs — Regulators require legal justifications. We translate local SHAP values into automated top-3 plain-English compliance summaries.",
        "Gap 5: Absence of Real-Time ML-BI Coupling — We deliver an integrated architecture connecting raw streaming ingestion, FastAPI inference, and Streamlit executive simulation."
    ]
    for g in gaps:
        p_g = tf.add_paragraph()
        p_g.text = f"• {g}"
        p_g.font.size = Pt(13)
    set_speaker_notes(s4, "Explain our five distinct innovations resolving the gaps identified in recent literature.")

    # Slide 5: System Architecture
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "End-to-End System Architecture & Execution Flow")
    if os.path.exists("docs/system_architecture.png"):
        s5.shapes.add_picture("docs/system_architecture.png", Inches(1.2), Inches(1.6), Inches(10.9), Inches(5.2))
    set_speaker_notes(s5, "Walk examiners through the modular architecture: raw data, zero-leakage prep, ensemble ML, SHAP transparency, BI layer, and deployed microservices.")

    # Slide 6: Dataset Description
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Dataset: Synthetic Anti-Money-Laundering (SAML-D)")
    tx = s6.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    ds_bullets = [
        "Benchmark Scale: 9,504,852 total financial records across 321 calendar days (Oct 2022 to Aug 2023).",
        "Severe Class Rarity: 9,873 confirmed laundering cases (0.104% positive prevalence).",
        "Key Operational Attributes: Time, Date, Sender_account, Receiver_account, Amount, Currencies, Locations, Payment Rails.",
        "Leakage Prevention Rule: 'Laundering_type' specifies illicit typology and is purged from features prior to training.",
        "Account IDs: Excluded from feature vector; leveraged strictly to compute causal behavioral history."
    ]
    for d in ds_bullets:
        p_d = tf.add_paragraph()
        p_d.text = f"• {d}"
        p_d.font.size = Pt(13)

    if os.path.exists("outputs/eda/01_class_imbalance.png"):
        s6.shapes.add_picture("outputs/eda/01_class_imbalance.png", Inches(7.2), Inches(1.8), Inches(5.3), Inches(4.8))
    set_speaker_notes(s6, "Highlight that SAML-D is an institutional-scale dataset of 9.5M rows with 0.104% positives, preserving real banking conditions.")

    # Slide 7: Preprocessing & Leakage Precautions
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Rigorous Preprocessing & Zero-Leakage Precautions")
    tx = s7.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    prep_pts = [
        "Chronological Partitioning: 70% Train (to May 18, 2023), 15% Validation (May 19 to Jul 5, 2023), 15% Untouched Test (Jul 6 to Aug 23, 2023).",
        "Verification Assertion: Automated check confirms max(train) < min(val) < min(test) with zero temporal leakage.",
        "Causal History Engineering (Shift-by-1): All customer transaction features use strictly earlier records only:",
        "   - sender_prior_count: Prior transaction frequency of sender.",
        "   - sender_prior_mean & sender_prior_std: Running historical financial mean and variance.",
        "   - amount_z_score: Dynamic amount deviation relative to customer's personal velocity.",
        "   - time_since_prev_sender_txn: Inter-transaction burst interval.",
        "   - distinct_prior_receivers: Fan-out counterparty proliferation metric.",
        "Encoders & Risk Jurisdictions: Fit strictly on TRAIN data; unseen values mapped to 'unknown'."
    ]
    for pt in prep_pts:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.size = Pt(13)
    set_speaker_notes(s7, "Emphasize how shift-by-1 causal features prevent look-ahead bias and explain why correlation checks confirmed no target leakage.")

    # Slide 8: EDA Highlights 1
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Exploratory Data Analysis: Risk Channels & Currencies")
    if os.path.exists("outputs/eda/03_laundering_rate_by_payment_type.png"):
        s8.shapes.add_picture("outputs/eda/03_laundering_rate_by_payment_type.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    if os.path.exists("outputs/eda/08_currency_mismatch_effect.png"):
        s8.shapes.add_picture("outputs/eda/08_currency_mismatch_effect.png", Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    set_speaker_notes(s8, "Point out the elevated risk in automated cash/cross-border payment channels and the multi-fold risk spike during currency conversion mismatch.")

    # Slide 9: EDA Highlights 2
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Exploratory Data Analysis: Geography & Typology Signatures")
    if os.path.exists("outputs/eda/12_location_pairs_heatmap.png"):
        s9.shapes.add_picture("outputs/eda/12_location_pairs_heatmap.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    if os.path.exists("outputs/eda/10_laundering_type_distribution.png"):
        s9.shapes.add_picture("outputs/eda/10_laundering_type_distribution.png", Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    set_speaker_notes(s9, "Highlight high-risk jurisdictional transfer corridors and the prevalence of structuring, fan-in, and single-large typologies.")

    # Slide 10: Model Building & Calibration
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Model Architecture: Rank-Calibrated Ensemble")
    tx = s10.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    m_pts = [
        "Baseline: Balanced Logistic Regression on standardized causal features (Val PR-AUC: 0.0313).",
        "M1 LightGBM: Gradient boosting with scale_pos_weight & early stopping (Val PR-AUC: 0.8883, Test PR-AUC: 0.9032).",
        "M2 Random Forest: Balanced subsampling on 1:10 downsampled training subset (Val PR-AUC: 0.8031).",
        "M3 Isolation Forest: Unsupervised subspace anomaly depth scoring.",
        "M4 Rank-Normalized Ensemble: Converts disparate model scores into empirical percentiles, optimized on validation PR-AUC:",
        "   Weights: 0.80 LGB + 0.15 RF + 0.025 IF + 0.025 LR.",
        "Isotonic Probability Calibration: Calibrates ensemble scores into defensible posterior probabilities."
    ]
    for pt in m_pts:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.size = Pt(12)

    if os.path.exists("outputs/figures/calibration_curve.png"):
        s10.shapes.add_picture("outputs/figures/calibration_curve.png", Inches(7.2), Inches(1.8), Inches(5.3), Inches(4.8))
    set_speaker_notes(s10, "Explain rank normalization for multi-model blending and how Isotonic Regression aligns predictions with empirical risk probabilities.")

    # Slide 11: Experimental Results
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, f"Results on Untouched TEST Split (PR-AUC: {m['pr_auc']:.4f})")
    if os.path.exists("outputs/figures/pr_curve.png"):
        s11.shapes.add_picture("outputs/figures/pr_curve.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    if os.path.exists("outputs/figures/confusion_matrix.png"):
        s11.shapes.add_picture("outputs/figures/confusion_matrix.png", Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    set_speaker_notes(s11, f"Report test results honestly: LightGBM achieves 0.9032 PR-AUC and ensemble achieves {m['pr_auc']:.4f}. At 1.0% review capacity, Recall@K is {m['recall_at_1.0pct']*100:.2f}%.")

    # Slide 12: Explainability & SHAP
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Explainability: SHAP Attribution & Regulatory Reasons")
    if os.path.exists("outputs/figures/shap_summary_bar.png"):
        s12.shapes.add_picture("outputs/figures/shap_summary_bar.png", Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    if os.path.exists("outputs/figures/shap_beeswarm.png"):
        s12.shapes.add_picture("outputs/figures/shap_beeswarm.png", Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8))
    set_speaker_notes(s12, "Explain that SHAP TreeExplainer attributes credit to account velocity and transaction amount, generating top-3 plain-English reasons for SAR reports.")

    # Slide 13: Business Risk & Cost Impact
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, f"Business Impact: ${bi['cost_saved_usd']/1e6:.2f}M Saved ({bi['pct_cost_saved']:.2f}% Cost Reduction)")
    if os.path.exists("outputs/figures/dashboard.png"):
        s13.shapes.add_picture("outputs/figures/dashboard.png", Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.0))
    set_speaker_notes(s13, f"Highlight the financial bottom line: our cost-optimal threshold reduces total loss from ${bi['cost_no_model_usd']/1e6:.2f}M to ${bi['cost_ensemble_usd']/1e6:.2f}M, saving ${bi['cost_saved_usd']/1e6:.2f}M.")

    # Slide 14: Robustness, Deployment & Conclusion
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Robustness, Deployment & Future Research Directions")
    tx = s14.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf = tx.text_frame
    tf.word_wrap = True
    c_pts = [
        "Imbalance Robustness: Sustained high ROC-AUC (0.9948) and PR-AUC across extreme negative-to-positive ratio variations (1:1 to 1000:1).",
        f"Adversarial Defense: Evasion attacks degraded naive recall, but adversarial retraining recovered detection to {adv_df.loc[adv_df['metric']=='Adversarially Retrained Recall (Defense Recovery)', 'value'].values[0]} with {adv_df.loc[adv_df['metric']=='Clean Test PR-AUC Post-Defense Retraining', 'value'].values[0]} clean PR-AUC.",
        "Production Deployment: Containerized FastAPI (Port 8000) and Streamlit (Port 7860) live microservices compatible with Hugging Face Spaces.",
        "Future Work: (1) Graph Neural Networks for multi-hop layering graphs, (2) Federated Learning across inter-bank consortiums, (3) Real-time Kafka/Flink streaming.",
        "Conclusion: A fully realized, reproducible ML-BI AML system achieving 93.58% Recall@1%, 97.94% cost reduction, and defensible audit explanations."
    ]
    for pt in c_pts:
        p_c = tf.add_paragraph()
        p_c.text = f"• {pt}"
        p_c.font.size = Pt(13)
    set_speaker_notes(s14, "Conclude with our stress-test results, deployment architecture, and key avenues for future research.")

    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
